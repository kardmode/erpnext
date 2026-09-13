# Copyright (c) 2020, Frappe Technologies Pvt. Ltd. and contributors
# For license information, please see license.txt


import copy
import io
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import frappe
from frappe import _
from frappe.desk.utils import provide_binary_file
from frappe.model.document import Document
from frappe.utils import add_days, add_months, flt, format_date, getdate, today
from frappe.utils.jinja import validate_template
from frappe.utils.pdf import get_pdf
from frappe.www.printview import get_print_style

from erpnext import get_company_currency
from erpnext.accounts.party import get_party_account_currency
from erpnext.accounts.report.accounts_receivable.accounts_receivable import execute as get_ar_soa
from erpnext.accounts.report.accounts_receivable_summary.accounts_receivable_summary import (
	execute as get_ageing,
)
from erpnext.accounts.report.general_ledger.general_ledger import execute as get_soa


class ProcessStatementOfAccounts(Document):
	def validate(self):
		if not self.subject:
			self.subject = "Statement Of Accounts for {{ customer.customer_name }}"
		from frappe.utils import strip_html

		if not self.body or not strip_html(self.body).strip():
			self.body = (
				"<p>To: <b>{{ customer.customer_name }}</b>,</p>"
				"<p>Please find attached your official <b>Statement of Accounts</b> from <b>{{ doc.company }}</b>"
				"{% if doc.from_date %} covering the period from <b>{{ doc.from_date }}</b> to <b>{{ doc.to_date }}</b>"
				"{% else %} as of <b>{{ doc.to_date or doc.posting_date }}</b>{% endif %}.</p>"
				"<p>Kindly review the statement and contact us should you have any questions or require any clarification.</p>"
				"<p>Thank you for your business.</p>"
				"<p>Kind regards,<br><b>{{ doc.company }}</b></p>"
			)
		if not self.pdf_name:
			self.pdf_name = "{{ customer.customer_name }}"

		validate_template(self.subject)
		validate_template(self.body)

		if not self.customers:
			frappe.throw(_("Customers not selected."))

		if self.enable_auto_email:
			if self.start_date and getdate(self.start_date) >= getdate(today()):
				self.to_date = self.start_date
				self.from_date = add_months(self.to_date, -1 * self.filter_duration)


def get_report_pdf(doc, consolidated=True):
	statement_dict = get_statement_dict(doc)
	if not bool(statement_dict):
		return False
	elif consolidated:
		delimiter = '<div style="page-break-before: always;"></div>' if doc.include_break else ""
		result = delimiter.join(list(statement_dict.values()))
		return get_pdf(result, {"orientation": doc.orientation})
	else:
		for customer, statement_html in statement_dict.items():
			statement_dict[customer] = get_pdf(statement_html, {"orientation": doc.orientation})
		return statement_dict


def get_customer_settlement_map(customer):
	"""
	Returns a dictionary containing:
	  - by_invoice: dict mapping target_invoice -> list of allocation dicts
	  - by_credit_note: dict mapping credit_note_name -> list of reallocation dicts
	"""
	allocations_by_inv = {}
	reallocations_by_cn = {}
	if not customer:
		return {"by_invoice": allocations_by_inv, "by_credit_note": reallocations_by_cn}

	from frappe.utils import flt

	# 1. Fetch all return / credit note Sales Invoices for this customer
	cn_invoices = set()
	try:
		sales_invoices = frappe.get_all(
			"Sales Invoice",
			filters={"customer": customer, "docstatus": 1},
			fields=["name", "is_return", "return_against", "grand_total"],
		)
		for si in sales_invoices:
			if si.is_return or flt(si.grand_total) < 0:
				cn_invoices.add(si.name)
	except Exception:
		pass

	# 2. Inspect Journal Entries for Credit Note / Reconciliation links
	jv_links = {}
	try:
		je_accounts = frappe.db.sql(
			"""
			SELECT 
				jea.parent as jv_name,
				jea.reference_name,
				jea.reference_type,
				jea.debit_in_account_currency as debit,
				jea.credit_in_account_currency as credit,
				je.voucher_type as jv_voucher_type,
				je.is_system_generated,
				je.user_remark,
				je.remark
			FROM `tabJournal Entry Account` jea
			JOIN `tabJournal Entry` je ON je.name = jea.parent
			WHERE je.docstatus = 1 AND jea.party = %s
			""",
			(customer,),
			as_dict=True,
		)
		jv_groups = {}
		for row in je_accounts:
			jv_groups.setdefault(row.jv_name, []).append(row)

		for jv_name, rows in jv_groups.items():
			source_cns = [r.reference_name for r in rows if r.reference_name and (r.reference_name in cn_invoices or "RET" in str(r.reference_name))]
			target_invs = [r.reference_name for r in rows if r.reference_name and r.reference_name not in cn_invoices and "RET" not in str(r.reference_name)]

			is_cn_jv = any(
				r.jv_voucher_type == "Credit Note"
				or "credit note" in str(r.user_remark or "").lower()
				or "credit note" in str(r.remark or "").lower()
				or (r.reference_name in cn_invoices)
			) or bool(source_cns)

			cn_ref = source_cns[0] if source_cns else None
			jv_links[jv_name] = {
				"is_credit_note": is_cn_jv,
				"credit_note_ref": cn_ref,
				"targets": target_invs,
			}

			if source_cns and target_invs:
				for src in source_cns:
					for tgt in target_invs:
						amt = min(
							sum(flt(r.debit or r.credit) for r in rows if r.reference_name == src),
							sum(flt(r.debit or r.credit) for r in rows if r.reference_name == tgt),
						)
						if amt > 0:
							reallocations_by_cn.setdefault(src, []).append(
								{"target_invoice": tgt, "amount": amt, "jv": jv_name}
							)
	except Exception:
		pass

	# 3. Check Payment Ledger Entry
	try:
		ple_rows = frappe.db.sql(
			"""
			SELECT voucher_no, against_voucher_no, amount, voucher_type, against_voucher_type
			FROM `tabPayment Ledger Entry`
			WHERE party = %s AND delinked = 0 AND against_voucher_no IS NOT NULL AND against_voucher_no != ''
			""",
			(customer,),
			as_dict=True,
		)
		ple_by_voucher = {}
		for ple in ple_rows:
			ple_by_voucher.setdefault(ple.voucher_no, []).append(ple)

		for ple in ple_rows:
			v_no = ple.voucher_no
			ag_no = ple.against_voucher_no
			amt = abs(flt(ple.amount))
			if amt <= 0 or v_no == ag_no:
				continue

			is_cn = False
			cn_ref = None

			if v_no in cn_invoices or "RET" in str(v_no) or ple.voucher_type == "Credit Note":
				is_cn = True
				cn_ref = v_no
			elif v_no in jv_links and jv_links[v_no]["is_credit_note"]:
				is_cn = True
				cn_ref = jv_links[v_no]["credit_note_ref"]
			else:
				peer_rows = ple_by_voucher.get(v_no, [])
				for peer in peer_rows:
					peer_ag = peer.against_voucher_no
					if peer_ag != ag_no and (peer_ag in cn_invoices or "RET" in str(peer_ag)):
						is_cn = True
						cn_ref = peer_ag
						break

			ref_no = ""
			mode_pay = ""
			if (ple.voucher_type == "Payment Entry") or ("PE-" in str(v_no)):
				pe_data = frappe.db.get_value("Payment Entry", v_no, ["reference_no", "mode_of_payment"], as_dict=True)
				if pe_data:
					ref_no = pe_data.reference_no or ""
					mode_pay = pe_data.mode_of_payment or ""
			elif (ple.voucher_type == "Journal Entry") or ("JV-" in str(v_no)):
				jv_data = frappe.db.get_value("Journal Entry", v_no, ["cheque_no", "mode_of_payment"], as_dict=True)
				if jv_data:
					ref_no = jv_data.cheque_no or ""
					mode_pay = jv_data.mode_of_payment or ""

			alloc_dict = {
				"voucher_no": v_no,
				"voucher_type": ple.voucher_type or ("Credit Note" if is_cn else "Journal Entry"),
				"is_credit_note": is_cn,
				"credit_note_ref": cn_ref,
				"is_advance": False,
				"amount": amt,
				"reference_no": ref_no,
				"mode_of_payment": mode_pay,
			}
			existing = allocations_by_inv.setdefault(ag_no, [])
			if not any(e["voucher_no"] == v_no and abs(e["amount"] - amt) < 0.01 for e in existing):
				existing.append(alloc_dict)

			if is_cn and cn_ref and cn_ref != ag_no:
				cn_reallocs = reallocations_by_cn.setdefault(cn_ref, [])
				if not any(r["target_invoice"] == ag_no and abs(r["amount"] - amt) < 0.01 for r in cn_reallocs):
					cn_reallocs.append({"target_invoice": ag_no, "amount": amt, "jv": v_no})

	except Exception:
		pass

	# 4. Check Sales Invoice Advance table
	try:
		adv_rows = frappe.db.sql(
			"""
			SELECT parent as target_invoice, reference_type, reference_name, allocated_amount as amount
			FROM `tabSales Invoice Advance`
			WHERE parent IN (SELECT name FROM `tabSales Invoice` WHERE customer = %s AND docstatus = 1)
			""",
			(customer,),
			as_dict=True,
		)
		for adv in adv_rows:
			amt = flt(adv.amount)
			if amt <= 0:
				continue
			tgt = adv.target_invoice
			ref_name = adv.reference_name
			ref_type = adv.reference_type
			is_cn = ref_name in cn_invoices or "RET" in str(ref_name) or ref_type in ["Credit Note"]
			if not is_cn and ref_name in jv_links:
				is_cn = jv_links[ref_name]["is_credit_note"]
				if is_cn and jv_links[ref_name]["credit_note_ref"]:
					ref_name = jv_links[ref_name]["credit_note_ref"]

			ref_no = ""
			mode_pay = ""
			if ref_type == "Payment Entry":
				pe_data = frappe.db.get_value("Payment Entry", ref_name, ["reference_no", "mode_of_payment"], as_dict=True)
				if pe_data:
					ref_no = pe_data.reference_no or ""
					mode_pay = pe_data.mode_of_payment or ""
			elif ref_type == "Journal Entry":
				jv_data = frappe.db.get_value("Journal Entry", ref_name, ["cheque_no", "mode_of_payment"], as_dict=True)
				if jv_data:
					ref_no = jv_data.cheque_no or ""
					mode_pay = jv_data.mode_of_payment or ""

			alloc_dict = {
				"voucher_no": ref_name,
				"voucher_type": ref_type,
				"is_credit_note": is_cn,
				"credit_note_ref": ref_name if is_cn else None,
				"is_advance": not is_cn,
				"amount": amt,
				"reference_no": ref_no,
				"mode_of_payment": mode_pay,
			}
			existing = allocations_by_inv.setdefault(tgt, [])
			if not any(e["voucher_no"] == ref_name and abs(e["amount"] - amt) < 0.01 for e in existing):
				existing.append(alloc_dict)

			if is_cn:
				cn_reallocs = reallocations_by_cn.setdefault(ref_name, [])
				if not any(r["target_invoice"] == tgt and abs(r["amount"] - amt) < 0.01 for r in cn_reallocs):
					cn_reallocs.append({"target_invoice": tgt, "amount": amt, "jv": ref_name})
	except Exception:
		pass

	# 5. Check Payment Entry Reference table
	try:
		pe_rows = frappe.db.sql(
			"""
			SELECT per.parent as pe_name, per.reference_name as target_invoice, per.allocated_amount as amount,
			       pe.payment_type, pe.reference_no, pe.mode_of_payment
			FROM `tabPayment Entry Reference` per
			JOIN `tabPayment Entry` pe ON pe.name = per.parent
			WHERE pe.docstatus = 1 AND pe.party = %s AND per.reference_doctype = 'Sales Invoice'
			""",
			(customer,),
			as_dict=True,
		)
		for pe in pe_rows:
			amt = flt(pe.amount)
			if amt <= 0:
				continue
			tgt = pe.target_invoice
			alloc_dict = {
				"voucher_no": pe.pe_name,
				"voucher_type": "Payment Entry",
				"is_credit_note": False,
				"credit_note_ref": None,
				"is_advance": (pe.payment_type == "Receive" and "ADV" in pe.pe_name),
				"amount": amt,
				"reference_no": pe.reference_no or "",
				"mode_of_payment": pe.mode_of_payment or "",
			}
			existing = allocations_by_inv.setdefault(tgt, [])
			if not any(e["voucher_no"] == pe.pe_name and abs(e["amount"] - amt) < 0.01 for e in existing):
				existing.append(alloc_dict)
	except Exception:
		pass

	return {"by_invoice": allocations_by_inv, "by_credit_note": reallocations_by_cn}


def enrich_statement_data(doc, entry, res):
	if not res or doc.report != "General Ledger":
		return res, {}

	settlement_data = get_customer_settlement_map(entry.customer)
	settlement_map = settlement_data["by_invoice"]
	cn_realloc_map = settlement_data["by_credit_note"]

	si_vouchers = []
	pe_vouchers = []
	jv_vouchers = []
	for row in res:
		if not isinstance(row, dict) or not row.get("posting_date"):
			continue
		v_type = row.get("voucher_type")
		v_no = row.get("voucher_no")
		if v_type == "Sales Invoice" and v_no:
			si_vouchers.append(v_no)
		elif v_type == "Payment Entry" and v_no:
			pe_vouchers.append(v_no)
		elif v_type == "Journal Entry" and v_no:
			jv_vouchers.append(v_no)

	si_meta = {}
	if si_vouchers:
		try:
			si_records = frappe.get_all(
				"Sales Invoice",
				filters={"name": ["in", list(set(si_vouchers))]},
				fields=["name", "is_return", "return_against", "project", "remarks"],
			)
			for si in si_records:
				si_meta[si.name] = si
		except Exception:
			pass

	pe_allocations = {}
	if pe_vouchers:
		try:
			pe_refs = frappe.get_all(
				"Payment Entry Reference",
				filters={"parent": ["in", list(set(pe_vouchers))], "allocated_amount": [">", 0]},
				fields=["parent", "reference_doctype", "reference_name", "allocated_amount"],
			)
			for ref in pe_refs:
				pe_allocations.setdefault(ref.parent, []).append(ref)
		except Exception:
			pass

	total_invoiced = 0.0
	total_credit_notes = 0.0
	total_payments = 0.0
	total_debit = 0.0
	total_credit = 0.0

	for row in res:
		if not isinstance(row, dict) or not row.get("posting_date"):
			continue
		v_type = row.get("voucher_type")
		v_no = row.get("voucher_no")
		debit = float(row.get("debit") or 0.0)
		credit = float(row.get("credit") or 0.0)

		total_debit += debit
		total_credit += credit

		if v_type == "Sales Invoice" and v_no in si_meta:
			meta = si_meta[v_no]
			row["invoice_remarks"] = meta.get("remarks")
			row["project"] = meta.get("project") or row.get("project")
			if meta.get("is_return") or credit > 0:
				row["is_return"] = 1
				row["return_against"] = meta.get("return_against")
				total_credit_notes += credit
				reallocs = cn_realloc_map.get(v_no, [])
				if reallocs:
					row["reconciled_allocations"] = reallocs
			else:
				total_invoiced += debit
				allocs = settlement_map.get(v_no, [])
				if allocs:
					row["reconciled_credits"] = allocs
		elif v_type == "Payment Entry":
			total_payments += credit
			if v_no in pe_allocations:
				row["allocations"] = pe_allocations[v_no]
		elif v_type == "Journal Entry":
			if credit > 0:
				total_payments += credit
			if debit > 0:
				total_invoiced += debit
		else:
			if debit > 0:
				total_invoiced += debit
			if credit > 0:
				total_payments += credit

	opening_bal = 0.0
	closing_bal = 0.0
	if len(res) > 0 and isinstance(res[0], dict):
		opening_bal = float(res[0].get("balance") or 0.0)
	if len(res) > 1 and isinstance(res[-1], dict):
		closing_bal = float(res[-1].get("balance") or 0.0)

	summary = {
		"opening_balance": opening_bal,
		"total_invoiced": total_invoiced,
		"total_credit_notes": total_credit_notes,
		"total_payments": total_payments,
		"total_debit": total_debit,
		"total_credit": total_credit,
		"closing_balance": closing_bal,
	}

	return res, summary


def filter_and_enrich_ar_data(doc, entry, res):
	if not res:
		return [], {}

	filtered_res = []
	selected_projects = [p.project_name for p in doc.project] if doc.project else []
	settlement_data = get_customer_settlement_map(entry.customer)
	settlement_map = settlement_data["by_invoice"]
	cn_realloc_map = settlement_data["by_credit_note"]

	# Get project, remarks, and return metadata for invoices
	si_list = [r.get("voucher_no") for r in res if isinstance(r, dict) and r.get("voucher_no")]
	si_projects = {}
	if si_list:
		try:
			si_meta = frappe.get_all(
				"Sales Invoice",
				filters={"name": ["in", list(set(si_list))]},
				fields=["name", "project", "remarks", "is_return", "return_against"],
			)
			for si in si_meta:
				si_projects[si.name] = si
		except Exception:
			pass

	# Get Payment Entry metadata (reference no, mode of payment, remarks, allocations)
	pe_meta = {}
	try:
		pe_names = [
			r.get("voucher_no")
			for r in res
			if isinstance(r, dict)
			and r.get("voucher_no")
			and (
				r.get("voucher_type") == "Payment Entry"
				or "PE-" in str(r.get("voucher_no"))
				or (float(r.get("invoiced") or 0.0) == 0 and float(r.get("outstanding") or 0.0) < 0)
			)
		]
		# Also collect any Payment Entry vouchers from the settlement map
		for alloc_list in settlement_map.values():
			for a in alloc_list:
				v = a.get("voucher_no")
				if v and ("PE-" in str(v) or a.get("voucher_type") == "Payment Entry"):
					pe_names.append(v)

		if pe_names:
			pe_records = frappe.get_all(
				"Payment Entry",
				filters={"name": ["in", list(set(pe_names))]},
				fields=["name", "reference_no", "reference_date", "mode_of_payment", "remarks", "received_amount", "paid_amount"],
			)
			for pe in pe_records:
				pe_meta[pe.name] = {
					"reference_no": pe.reference_no,
					"reference_date": pe.reference_date,
					"mode_of_payment": pe.mode_of_payment,
					"remarks": pe.remarks,
					"total_amount": flt(pe.received_amount or pe.paid_amount),
					"settled_invoices": [],
				}

			pe_refs = frappe.db.sql(
				"""
				SELECT parent, reference_name, allocated_amount
				FROM `tabPayment Entry Reference`
				WHERE parent IN %(pe_names)s AND reference_doctype = 'Sales Invoice' AND allocated_amount > 0
				""",
				{"pe_names": list(set(pe_names))},
				as_dict=True,
			)
			for pr in pe_refs:
				if pr.parent in pe_meta:
					pe_meta[pr.parent]["settled_invoices"].append(
						{"invoice": pr.reference_name, "amount": flt(pr.allocated_amount)}
					)
	except Exception:
		pass

	# Also enrich settlement_map allocations with payment metadata (reference_no, mode_of_payment)
	if pe_meta:
		for alloc_list in settlement_map.values():
			for a in alloc_list:
				v = a.get("voucher_no")
				if v in pe_meta:
					a["reference_no"] = pe_meta[v].get("reference_no")
					a["mode_of_payment"] = pe_meta[v].get("mode_of_payment")

	total_gross_invoiced = 0.0
	total_paid = 0.0
	total_credit_notes_issued = 0.0
	total_outstanding = 0.0

	for row in res:
		if not isinstance(row, dict):
			continue
		v_no = row.get("voucher_no")
		v_date = row.get("posting_date")
		si_info = si_projects.get(v_no, {})

		if si_info:
			row["project"] = si_info.get("project") or row.get("project")
			row["invoice_remarks"] = si_info.get("remarks")
			row["is_return"] = si_info.get("is_return")
			row["return_against"] = si_info.get("return_against")

		# Apply project filter
		if selected_projects and row.get("project") not in selected_projects:
			continue

		# Apply from_date filter if posting_date exists
		if doc.from_date and v_date and getdate(v_date) < getdate(doc.from_date):
			continue

		# Apply to_date filter if posting_date exists
		if doc.to_date and v_date and getdate(v_date) > getdate(doc.to_date):
			continue

		invoiced = float(row.get("invoiced") or row.get("invoiced_amount") or 0.0)
		paid = float(row.get("paid") or row.get("paid_amount") or 0.0)
		credit_note = float(row.get("credit_note") or row.get("credit_note_amount") or 0.0)
		outstanding = float(row.get("outstanding") or row.get("outstanding_amount") or 0.0)

		# Check if this row is a standalone Payment Entry or advance credit
		is_pe = (
			row.get("voucher_type") == "Payment Entry"
			or (v_no and "PE-" in str(v_no))
			or (invoiced == 0 and outstanding < 0)
			or (v_no in pe_meta)
		)
		if is_pe:
			row["is_payment_entry"] = 1
			row["status"] = "Unallocated Payment"
			if v_no in pe_meta:
				p_data = pe_meta[v_no]
				row["customer_ref"] = p_data.get("reference_no")
				row["reference_date"] = p_data.get("reference_date")
				row["mode_of_payment"] = p_data.get("mode_of_payment")
				row["pe_remarks"] = p_data.get("remarks")
				row["pe_total_amount"] = p_data.get("total_amount")
				row["pe_settled_invoices"] = p_data.get("settled_invoices")

			total_paid += paid
			total_outstanding += outstanding
		elif row.get("is_return") or invoiced < 0:
			row["is_return"] = 1
			reallocs = cn_realloc_map.get(v_no, [])
			row["reallocated_to"] = reallocs
			abs_cn = abs(invoiced) if invoiced != 0 else abs(credit_note)
			realloc_sum = sum(float(r.get("amount", 0.0)) for r in reallocs)
			unallocated_cn = max(0.0, abs_cn - realloc_sum)

			row["invoiced"] = -abs_cn
			row["paid"] = 0.0
			row["credit_note"] = 0.0
			row["outstanding"] = -unallocated_cn if unallocated_cn > 0 else 0.0
			row["status"] = "Credit Note"
			total_credit_notes_issued += abs_cn
		else:
			allocs = settlement_map.get(v_no, [])
			if allocs:
				row["allocations"] = allocs

			cn_alloc_total = sum(a["amount"] for a in allocs if a.get("is_credit_note"))
			paid_alloc_total = sum(a["amount"] for a in allocs if not a.get("is_credit_note"))
			if cn_alloc_total > 0 and credit_note == 0:
				credit_note = cn_alloc_total
				if paid >= cn_alloc_total:
					paid = max(0.0, paid - cn_alloc_total)
				row["credit_note"] = credit_note
				row["paid"] = paid

			if outstanding == 0:
				if credit_note > 0 and paid == 0:
					row["status"] = "Paid via Credit Note"
				else:
					row["status"] = "Paid in Full"
			elif outstanding < invoiced:
				row["status"] = "Partially Paid"
			else:
				row["status"] = "Unpaid"

			total_gross_invoiced += invoiced
			total_paid += paid
			total_outstanding += outstanding

		filtered_res.append(row)

	total_net_invoiced = max(0.0, total_gross_invoiced - total_credit_notes_issued)

	summary = {
		"opening_balance": 0.0,
		"gross_invoiced": total_gross_invoiced,
		"total_credit_notes": total_credit_notes_issued,
		"net_invoiced": total_net_invoiced,
		"total_invoiced": total_gross_invoiced,
		"total_paid": total_paid,
		"total_payments": total_paid,
		"closing_balance": total_outstanding,
		"total_outstanding": total_outstanding,
	}

	return filtered_res, summary


def get_project_billing_data(doc, entry):
	filters = {
		"customer": entry.customer,
		"docstatus": 1,
	}
	if doc.project:
		filters["project"] = ["in", [p.project_name for p in doc.project]]
	if doc.from_date:
		filters["posting_date"] = [">=", doc.from_date]
	if doc.to_date or doc.posting_date:
		to_d = doc.to_date or doc.posting_date
		if doc.from_date:
			filters["posting_date"] = ["between", [doc.from_date, to_d]]
		else:
			filters["posting_date"] = ["<=", to_d]

	invoices = frappe.get_all(
		"Sales Invoice",
		filters=filters,
		fields=[
			"name",
			"posting_date",
			"due_date",
			"grand_total",
			"outstanding_amount",
			"project",
			"is_return",
			"return_against",
			"remarks",
			"currency",
		],
		order_by="posting_date asc, creation asc",
	)

	if not invoices:
		return [], {}, {}

	from frappe.utils import flt
	settlement_data = get_customer_settlement_map(entry.customer)
	settlement_map = settlement_data["by_invoice"]
	cn_realloc_map = settlement_data["by_credit_note"]

	res = []
	total_gross_invoiced = 0.0
	total_paid = 0.0
	total_credit_notes_issued = 0.0
	total_outstanding = 0.0

	for inv in invoices:
		invoiced_amt = flt(inv.grand_total)
		out_amt = flt(inv.outstanding_amount)

		if inv.is_return or invoiced_amt < 0:
			abs_cn = abs(invoiced_amt)
			reallocs = cn_realloc_map.get(inv.name, [])
			total_credit_notes_issued += abs_cn
			realloc_sum = sum(float(r.get("amount", 0.0)) for r in reallocs)
			unallocated_cn = max(0.0, abs_cn - realloc_sum)

			row = {
				"posting_date": inv.posting_date,
				"due_date": inv.due_date,
				"voucher_no": inv.name,
				"voucher_type": "Sales Invoice",
				"project": inv.project,
				"is_return": 1,
				"return_against": inv.return_against,
				"remarks": inv.remarks,
				"invoice_remarks": inv.remarks,
				"invoiced": -abs_cn,
				"paid": 0.0,
				"credit_note": 0.0,
				"outstanding": -unallocated_cn if unallocated_cn > 0 else 0.0,
				"reallocated_to": reallocs,
				"allocations": [],
				"status": "Credit Note",
			}
			res.append(row)
		else:
			allocs = settlement_map.get(inv.name, [])
			cr_amt = sum(a["amount"] for a in allocs if a.get("is_credit_note"))
			paid_amt = sum(a["amount"] for a in allocs if not a.get("is_credit_note"))

			settled_diff = invoiced_amt - out_amt
			if settled_diff > 0 and (paid_amt + cr_amt) == 0:
				paid_amt = settled_diff

			if out_amt == 0:
				if cr_amt > 0 and paid_amt == 0:
					status_label = "Paid via Credit Note"
				else:
					status_label = "Paid in Full"
			elif out_amt < invoiced_amt:
				status_label = "Partially Paid"
			else:
				status_label = "Unpaid"

			total_gross_invoiced += invoiced_amt
			total_paid += paid_amt
			total_outstanding += out_amt

			row = {
				"posting_date": inv.posting_date,
				"due_date": inv.due_date,
				"voucher_no": inv.name,
				"voucher_type": "Sales Invoice",
				"project": inv.project,
				"is_return": 0,
				"return_against": inv.return_against,
				"remarks": inv.remarks,
				"invoice_remarks": inv.remarks,
				"invoiced": invoiced_amt,
				"paid": paid_amt,
				"credit_note": cr_amt,
				"outstanding": out_amt,
				"allocations": allocs,
				"status": status_label,
			}
			res.append(row)

	total_net_invoiced = max(0.0, total_gross_invoiced - total_credit_notes_issued)

	summary = {
		"opening_balance": 0.0,
		"gross_invoiced": total_gross_invoiced,
		"total_credit_notes": total_credit_notes_issued,
		"net_invoiced": total_net_invoiced,
		"total_invoiced": total_gross_invoiced,
		"total_paid": total_paid,
		"total_credit_notes": total_credit_notes_issued,
		"total_payments": total_paid,
		"closing_balance": total_outstanding,
		"total_outstanding": total_outstanding,
	}

	# Build Payments list for Mode 3
	invoice_names = [inv.name for inv in invoices if not (inv.is_return or flt(inv.grand_total) < 0)]
	payments_map = {}
	if invoice_names:
		try:
			pe_query_rows = frappe.db.sql(
				"""
				SELECT pe.name as voucher_no, pe.posting_date, pe.remarks, pe.reference_no, pe.mode_of_payment, per.allocated_amount as amount, per.reference_name as invoice_no
				FROM `tabPayment Entry Reference` per
				JOIN `tabPayment Entry` pe ON pe.name = per.parent
				WHERE pe.docstatus = 1 AND per.reference_doctype = 'Sales Invoice' AND per.reference_name IN %(inv_names)s
				ORDER BY pe.posting_date asc, pe.creation asc
				""",
				{"inv_names": invoice_names},
				as_dict=True,
			)
			for r in pe_query_rows:
				vno = r.voucher_no
				if vno not in payments_map:
					payments_map[vno] = {
						"voucher_no": vno,
						"posting_date": r.posting_date,
						"remarks": r.remarks,
						"reference_no": r.reference_no,
						"mode_of_payment": r.mode_of_payment,
						"amount": 0.0,
						"invoices": [],
					}
				payments_map[vno]["amount"] += flt(r.amount)
				if r.invoice_no and r.invoice_no not in payments_map[vno]["invoices"]:
					payments_map[vno]["invoices"].append(r.invoice_no)
		except Exception:
			pass

	# Also check if any JV settlement or other payment in allocations
	for inv_name in invoice_names:
		for alloc in settlement_map.get(inv_name, []):
			if not alloc.get("is_credit_note"):
				vno = alloc.get("voucher_no")
				if vno and vno not in payments_map:
					payments_map[vno] = {
						"voucher_no": vno,
						"posting_date": None,
						"remarks": "Settlement Adjustment / JV",
						"amount": flt(alloc.get("amount", 0.0)),
						"invoices": [inv_name],
					}

	payments_list = list(payments_map.values())
	if not payments_list and total_paid > 0:
		payments_list.append({
			"voucher_no": "Payments Received",
			"posting_date": None,
			"remarks": "Recorded payments and receipts against project invoices",
			"amount": total_paid,
			"invoices": invoice_names,
		})

	# Build Credit Notes list for Mode 3
	credit_notes_list = []
	for inv in invoices:
		invoiced_amt = flt(inv.grand_total)
		if inv.is_return or invoiced_amt < 0:
			abs_cn = abs(invoiced_amt)
			reallocs = cn_realloc_map.get(inv.name, [])
			credit_notes_list.append({
				"voucher_no": inv.name,
				"return_against": inv.return_against,
				"amount": abs_cn,
				"remarks": inv.remarks,
				"allocations": reallocs,
			})

	# Build Outstanding Invoices list for Mode 3
	outstanding_list = []
	for inv in invoices:
		invoiced_amt = flt(inv.grand_total)
		out_amt = flt(inv.outstanding_amount)
		if not (inv.is_return or invoiced_amt < 0) and out_amt > 0.001:
			outstanding_list.append({
				"posting_date": inv.posting_date,
				"due_date": inv.due_date,
				"voucher_no": inv.name,
				"invoiced": invoiced_amt,
				"outstanding": out_amt,
				"paid": invoiced_amt - out_amt,
			})

	rec_data = {
		"payments": payments_list,
		"credit_notes": credit_notes_list,
		"outstanding_invoices": outstanding_list,
		"all_invoices": invoices,
	}

	return res, summary, rec_data


def get_statement_dict(doc, get_statement_dict=False):
	statement_dict = {}
	ageing = ""

	for entry in doc.customers:
		if doc.include_ageing:
			ageing = set_ageing(doc, entry)

		tax_id = frappe.get_doc("Customer", entry.customer).tax_id
		presentation_currency = (
			get_party_account_currency("Customer", entry.customer, doc.company)
			or doc.currency
			or get_company_currency(doc.company)
		)

		filters = get_common_filters(doc)
		filters["presentation_currency"] = presentation_currency

		if doc.ignore_exchange_rate_revaluation_journals:
			filters.update({"ignore_err": True})

		if doc.ignore_cr_dr_notes:
			filters.update({"ignore_cr_dr_notes": True})

		summary = {}
		rec_data = {}
		if doc.report == "General Ledger":
			filters.update(get_gl_filters(doc, entry, tax_id, presentation_currency))
			col, res = get_soa(filters)
			for x in [0, -2, -1]:
				if len(res) > abs(x) and isinstance(res[x], dict) and "account" in res[x]:
					res[x]["account"] = str(res[x]["account"]).replace("'", "")
			res, summary = enrich_statement_data(doc, entry, res)
		elif doc.report == "Project Billing Statement" or (
			doc.report == "Accounts Receivable" and doc.get("include_settled_invoices")
		):
			filters.update(get_ar_filters(doc, entry))
			col = []
			res, summary, rec_data = get_project_billing_data(doc, entry)
			if not res:
				continue
		else:
			filters.update(get_ar_filters(doc, entry))
			ar_res = get_ar_soa(filters)
			col, raw_res = ar_res[0], ar_res[1]
			if not raw_res:
				continue
			res, summary = filter_and_enrich_ar_data(doc, entry, raw_res)
			if not res:
				continue

		statement_dict[entry.customer] = (
			[res, ageing, summary]
			if get_statement_dict
			else get_html(doc, filters, entry, col, res, ageing, summary, rec_data)
		)

	return statement_dict


def set_ageing(doc, entry):
	report_date = doc.posting_date or doc.to_date or today()
	ageing_filters = frappe._dict(
		{
			"company": doc.company,
			"report_date": report_date,
			"ageing_based_on": doc.ageing_based_on,
			"range1": 30,
			"range2": 60,
			"range3": 90,
			"range4": 120,
			"party_type": "Customer",
			"party": [entry.customer],
		}
	)
	col1, ageing = get_ageing(ageing_filters)

	if ageing:
		ageing[0]["ageing_based_on"] = doc.ageing_based_on

	return ageing


def get_common_filters(doc):
	return frappe._dict(
		{
			"company": doc.company,
			"finance_book": doc.finance_book if doc.finance_book else None,
			"account": [doc.account] if doc.account else None,
			"cost_center": [cc.cost_center_name for cc in doc.cost_center] if doc.cost_center else None,
			"show_remarks": doc.show_remarks,
		}
	)


def get_gl_filters(doc, entry, tax_id, presentation_currency):
	to_d = getdate(doc.to_date or doc.posting_date or today())
	from_d = getdate(doc.from_date or "1900-01-01")
	return {
		"company": doc.company,
		"from_date": from_d,
		"to_date": to_d,
		"party_type": "Customer",
		"party": [entry.customer],
		"party_name": [entry.customer_name] if entry.customer_name else None,
		"presentation_currency": presentation_currency,
		"categorize_by": doc.categorize_by or "Categorize by Voucher",
		"currency": doc.currency,
		"project": [p.project_name for p in doc.project] if doc.project else None,
		"show_opening_entries": 0,
		"include_default_book_entries": 0,
		"tax_id": tax_id if tax_id else None,
		"show_net_values_in_party_account": doc.show_net_values_in_party_account,
	}


def get_ar_filters(doc, entry):
	report_date = getdate(doc.posting_date or doc.to_date or today())
	from_d = getdate(doc.from_date) if doc.from_date else None
	to_d = getdate(doc.to_date) if doc.to_date else report_date
	filters = {
		"report_date": report_date,
		"from_date": from_d,
		"to_date": to_d,
		"party_type": "Customer",
		"party": [entry.customer],
		"customer_name": entry.customer_name if entry.customer_name else None,
		"payment_terms_template": doc.payment_terms_template if doc.payment_terms_template else None,
		"sales_partner": doc.sales_partner if doc.sales_partner else None,
		"sales_person": doc.sales_person if doc.sales_person else None,
		"territory": doc.territory if doc.territory else None,
		"based_on_payment_terms": doc.based_on_payment_terms,
		"report_name": "Accounts Receivable",
		"ageing_based_on": doc.ageing_based_on,
		"range1": 30,
		"range2": 60,
		"range3": 90,
		"range4": 120,
	}
	if doc.project:
		filters["project"] = [p.project_name for p in doc.project]
	if doc.cost_center:
		filters["cost_center"] = [cc.cost_center_name for cc in doc.cost_center]
	return filters


def get_html(doc, filters, entry, col, res, ageing, summary=None, rec_data=None):
	base_template_path = "frappe/www/printview.html"
	if doc.report == "General Ledger":
		template_path = (
			"erpnext/accounts/doctype/process_statement_of_accounts/process_statement_of_accounts.html"
		)
	elif doc.report == "Project Billing Statement":
		template_path = "erpnext/accounts/doctype/process_statement_of_accounts/process_statement_of_accounts_project_reconciliation.html"
	else:
		template_path = "erpnext/accounts/doctype/process_statement_of_accounts/process_statement_of_accounts_accounts_receivable.html"

	process_soa_html = frappe.get_hooks("process_soa_html")
	# fetching custom print format for Process Statement of Accounts
	if process_soa_html and process_soa_html.get(doc.report):
		template_path = process_soa_html[doc.report][-1]

	letter_head = None
	if doc.letter_head:
		from frappe.www.printview import get_letter_head

		letter_head = get_letter_head(doc, 0)
	html = frappe.render_template(
		template_path,
		{
			"filters": filters,
			"currency": filters.get("presentation_currency"),
			"data": res,
			"report": {"report_name": doc.report, "columns": col},
			"ageing": ageing[0] if (doc.include_ageing and ageing) else None,
			"letter_head": letter_head if doc.letter_head else None,
			"summary": summary or {},
			"rec_data": rec_data or {},
			"customer_doc": frappe.get_cached_doc("Customer", entry.customer),
			"terms_and_conditions": frappe.db.get_value(
				"Terms and Conditions", doc.terms_and_conditions, "terms"
			)
			if doc.terms_and_conditions
			else None,
		},
	)
	html = frappe.render_template(
		base_template_path,
		{"body": html, "css": get_print_style(), "title": "Statement For " + entry.customer},
	)
	return html


def get_customers_based_on_territory_or_customer_group(customer_collection, collection_name):
	fields_dict = {
		"Customer Group": "customer_group",
		"Territory": "territory",
	}
	collection = frappe.get_doc(customer_collection, collection_name)
	selected = [
		customer.name
		for customer in frappe.get_list(
			customer_collection,
			filters=[["lft", ">=", collection.lft], ["rgt", "<=", collection.rgt]],
			fields=["name"],
			order_by="lft asc, rgt desc",
		)
	]
	return frappe.get_list(
		"Customer",
		fields=["name", "customer_name", "email_id"],
		filters=[[fields_dict[customer_collection], "IN", selected]],
	)


def get_customers_based_on_sales_person(sales_person):
	lft, rgt = frappe.db.get_value("Sales Person", sales_person, ["lft", "rgt"])
	records = frappe.db.sql(
		"""
		select distinct parent, parenttype
		from `tabSales Team` steam
		where parenttype = 'Customer'
			and exists(select name from `tabSales Person` where lft >= %s and rgt <= %s and name = steam.sales_person)
	""",
		(lft, rgt),
		as_dict=1,
	)
	sales_person_records = frappe._dict()
	for d in records:
		sales_person_records.setdefault(d.parenttype, set()).add(d.parent)
	if sales_person_records.get("Customer"):
		return frappe.get_list(
			"Customer",
			fields=["name", "customer_name", "email_id"],
			filters=[["name", "in", list(sales_person_records["Customer"])]],
		)
	else:
		return []


def get_recipients_and_cc(customer, doc):
	# If Override Recipient is set, send ONLY to the override address(es) and suppress customer contacts
	if doc.get("override_email") and doc.override_email.strip():
		override_list = [e.strip() for e in doc.override_email.split(",") if e.strip()]
		return override_list, []

	recipients = []
	for clist in doc.customers:
		if clist.customer == customer:
			if clist.billing_email:
				for email in clist.billing_email.split(","):
					if email.strip():
						recipients.append(email.strip())
			if doc.primary_mandatory and clist.primary_email:
				for email in clist.primary_email.split(","):
					if email.strip():
						recipients.append(email.strip())

	cc = []
	if doc.cc_to != "":
		try:
			cc_email = frappe.get_value("User", doc.cc_to, "email")
			if cc_email:
				# If no customer recipients found, fallback cc_to into primary recipients so email can send
				if not recipients:
					recipients.append(cc_email)
				else:
					cc.append(cc_email)
		except Exception:
			pass

	return recipients, cc


def get_context(customer, doc):
	template_doc = copy.deepcopy(doc)
	del template_doc.customers
	if template_doc.from_date:
		template_doc.from_date = format_date(template_doc.from_date)
	if template_doc.to_date:
		template_doc.to_date = format_date(template_doc.to_date)
	return {
		"doc": template_doc,
		"customer": frappe.get_doc("Customer", customer),
		"frappe": frappe.utils,
	}


@frappe.whitelist()
def fetch_customers(customer_collection, collection_name, primary_mandatory):
	customer_list = []
	customers = []

	if customer_collection == "Sales Person":
		customers = get_customers_based_on_sales_person(collection_name)
		if not bool(customers):
			frappe.throw(_("No Customers found with selected options."))
	else:
		if customer_collection == "Sales Partner":
			customers = frappe.get_list(
				"Customer",
				fields=["name", "customer_name", "email_id"],
				filters=[["default_sales_partner", "=", collection_name]],
			)
		else:
			customers = get_customers_based_on_territory_or_customer_group(
				customer_collection, collection_name
			)

	for customer in customers:
		primary_email = customer.get("email_id") or ""
		billing_email = get_customer_emails(customer.name, 1, billing_and_primary=False)

		if int(primary_mandatory):
			if primary_email == "":
				continue
		elif (billing_email == "") and (primary_email == ""):
			continue

		customer_list.append(
			{
				"name": customer.name,
				"customer_name": customer.customer_name,
				"primary_email": primary_email,
				"billing_email": billing_email,
			}
		)
	return customer_list


@frappe.whitelist()
def get_customer_emails(customer_name, primary_mandatory, billing_and_primary=True):
	"""Returns first email from Contact Email table as a Billing email
	when Is Billing Contact checked
	and Primary email- email with Is Primary checked"""

	billing_email = frappe.db.sql(
		"""
		SELECT
			email.email_id
		FROM
			`tabContact Email` AS email
		JOIN
			`tabDynamic Link` AS link
		ON
			email.parent=link.parent
		JOIN
			`tabContact` AS contact
		ON
			contact.name=link.parent
		WHERE
			link.link_doctype='Customer'
			and link.link_name=%s
			and contact.is_billing_contact=1
		ORDER BY
			contact.creation desc""",
		customer_name,
	)

	if len(billing_email) == 0 or (billing_email[0][0] is None):
		if billing_and_primary:
			frappe.throw(_("No billing email found for customer: {0}").format(customer_name))
		else:
			return ""

	if billing_and_primary:
		primary_email = frappe.get_value("Customer", customer_name, "email_id")
		if primary_email is None and int(primary_mandatory):
			frappe.throw(_("No primary email found for customer: {0}").format(customer_name))
		return [primary_email or "", billing_email[0][0]]
	else:
		return billing_email[0][0] or ""


def get_statement_download_filename(doc, extension="pdf"):
	try:
		mode_slug = re.sub(r"[^\w\-_]", "_", str(doc.report or "Statement")).strip("_")
		parts = []
		if doc.get("customers") and len(doc.customers) == 1:
			cust = str(doc.customers[0].get("customer_name") or doc.customers[0].get("customer") or "")
			if cust:
				parts.append(re.sub(r"[^\w\-_]", "_", cust).strip("_"))
		if not parts and doc.name:
			parts.append(re.sub(r"[^\w\-_]", "_", str(doc.name)).strip("_"))

		parts.append(mode_slug)
		clean_name = "_".join([p for p in parts if p])
		return clean_name if extension == "" else f"{clean_name}.{extension}"
	except Exception:
		return f"{doc.name or 'statement'}.{extension}" if extension else (doc.name or "statement")


@frappe.whitelist()
def download_statements(document_name):
	try:
		doc = frappe.get_doc("Process Statement Of Accounts", document_name)
		report = get_report_pdf(doc)
		if report:
			frappe.local.response.filename = get_statement_download_filename(doc, "pdf")
			frappe.local.response.filecontent = report
			frappe.local.response.type = "download"
		else:
			frappe.msgprint(_("No data found for the selected criteria."))
	except Exception as e:
		frappe.log_error(f"Error in download_statements: {str(e)}", "Process Statement Of Accounts")
		frappe.throw(str(e))


@frappe.whitelist()
def download_excel_statements(document_name):
	try:
		doc = frappe.get_doc("Process Statement Of Accounts", document_name)
		excel_file = get_report_excel(doc)
		if excel_file:
			filename_base = get_statement_download_filename(doc, "")
			provide_binary_file(filename_base, "xlsx", excel_file.getvalue())
		else:
			frappe.msgprint(_("No data found for the selected criteria."))
	except Exception as e:
		frappe.log_error(f"Error in download_excel_statements: {str(e)}", "Process Statement Of Accounts")
		frappe.throw(str(e))


def get_report_excel(doc):
	wb = openpyxl.Workbook()
	wb.remove(wb.active)  # remove default sheet

	thin_border = Border(
		left=Side(style="thin", color="CBD5E1"),
		right=Side(style="thin", color="CBD5E1"),
		top=Side(style="thin", color="CBD5E1"),
		bottom=Side(style="thin", color="CBD5E1"),
	)
	header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
	sub_header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
	card_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
	total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

	header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
	title_font = Font(name="Calibri", size=15, bold=True, color="0F172A")
	section_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
	bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
	regular_font = Font(name="Calibri", size=10, color="1E293B")

	sheets_created = 0

	for entry in doc.customers:
		ageing = ""
		if doc.include_ageing:
			ageing = set_ageing(doc, entry)

		tax_id = frappe.get_doc("Customer", entry.customer).tax_id
		presentation_currency = (
			get_party_account_currency("Customer", entry.customer, doc.company)
			or doc.currency
			or get_company_currency(doc.company)
		)

		filters = get_common_filters(doc)
		if doc.ignore_exchange_rate_revaluation_journals:
			filters.update({"ignore_err": True})
		if doc.ignore_cr_dr_notes:
			filters.update({"ignore_cr_dr_notes": True})

		summary = {}
		rec_data = {}
		if doc.report == "General Ledger":
			filters.update(get_gl_filters(doc, entry, tax_id, presentation_currency))
			col, res = get_soa(filters)
			for x in [0, -2, -1]:
				if len(res) > abs(x) and isinstance(res[x], dict) and "account" in res[x]:
					res[x]["account"] = str(res[x]["account"]).replace("'", "")
			res, summary = enrich_statement_data(doc, entry, res)
		elif doc.report == "Project Billing Statement" or (
			doc.report == "Accounts Receivable" and doc.get("include_settled_invoices")
		):
			filters.update(get_ar_filters(doc, entry))
			res, summary, rec_data = get_project_billing_data(doc, entry)
			if not res:
				continue
		else:
			filters.update(get_ar_filters(doc, entry))
			ar_res = get_ar_soa(filters)
			col, raw_res = ar_res[0], ar_res[1]
			if not raw_res:
				continue
			res, summary = filter_and_enrich_ar_data(doc, entry, raw_res)
			if not res:
				continue

		sheet_title = re.sub(r"[\\/*?:\[\]]", "_", entry.customer_name or entry.customer)[:30]
		ws = wb.create_sheet(title=sheet_title)
		sheets_created += 1

		# Header block
		ws["A1"] = f"{doc.company} - Statement of Accounts ({doc.report})"
		ws["A1"].font = title_font
		ws["A2"] = f"Customer: {entry.customer_name or entry.customer} ({entry.customer})"
		ws["A2"].font = section_font

		from_d = filters.get("from_date") or doc.from_date
		to_d = filters.get("to_date") or doc.to_date or filters.get("report_date")
		if from_d:
			period_str = f"Period: {frappe.format(from_d, 'Date')} to {frappe.format(to_d, 'Date')}"
		else:
			period_str = f"Until: {frappe.format(to_d, 'Date')}"

		ws["A3"] = f"{period_str} | Currency: {presentation_currency} | Tax ID: {tax_id or 'N/A'}"
		ws["A3"].font = regular_font

		current_row = 5

		if doc.report == "Project Billing Statement":
			gross_inv = summary.get("gross_invoiced") or summary.get("total_invoiced", 0.0)
			cn_amt = summary.get("total_credit_notes", 0.0)
			net_inv = summary.get("net_invoiced") or (gross_inv - cn_amt)
			paid_amt = summary.get("total_paid", 0.0)
			out_amt = summary.get("total_outstanding", 0.0)

			# Executive KPI Summary Cards
			ws.cell(row=current_row, column=1, value="EXECUTIVE SUMMARY").font = section_font
			current_row += 1

			kpis = [
				("Total Invoiced (Gross)", gross_inv),
				("Credit Notes & Returns", -cn_amt if cn_amt else 0.0),
				("Net Billed Amount", net_inv),
				("Total Paid Amount", paid_amt),
				("Net Outstanding Due", out_amt),
			]

			for col_idx, (label, val) in enumerate(kpis, start=1):
				lbl_cell = ws.cell(row=current_row, column=col_idx, value=label)
				lbl_cell.font = bold_font
				lbl_cell.fill = sub_header_fill
				lbl_cell.border = thin_border
				lbl_cell.alignment = Alignment(horizontal="center")

				val_cell = ws.cell(row=current_row + 1, column=col_idx, value=float(val or 0.0))
				val_cell.font = bold_font
				val_cell.number_format = "#,##0.00"
				val_cell.fill = card_fill
				val_cell.border = thin_border
				val_cell.alignment = Alignment(horizontal="right")

			current_row += 3

			# 1. Project Billing & Payment Summary
			ws.cell(row=current_row, column=1, value="1. PROJECT BILLING & PAYMENT SUMMARY").font = section_font
			current_row += 1

			sec1_headers = ["Description / Commercial Stage", "Amount"]
			for col_idx, h in enumerate(sec1_headers, start=1):
				c = ws.cell(row=current_row, column=col_idx, value=h)
				c.font = header_font
				c.fill = header_fill
				c.border = thin_border
				c.alignment = Alignment(horizontal="center" if col_idx == 1 else "right")
			current_row += 1

			gross_inv = summary.get("gross_invoiced") or summary.get("total_invoiced", 0.0)
			cn_amt = summary.get("total_credit_notes", 0.0)
			net_inv = summary.get("net_invoiced") or (gross_inv - cn_amt)
			paid_amt = summary.get("total_paid", 0.0)
			out_amt = summary.get("total_outstanding", 0.0)

			summary_rows = [
				("Total Invoices Issued (Gross Billing)", gross_inv, False, False),
				("Less: Credit Note Adjustments", -cn_amt if cn_amt else 0.0, False, False),
				("Net Invoiced Amount (Net Project Value)", net_inv, True, False),
				("Less: Payments Received", -paid_amt if paid_amt else 0.0, False, False),
				("NET BALANCE OUTSTANDING DUE", out_amt, True, True),
			]

			for desc, val, is_subtotal, is_grand_total in summary_rows:
				c_desc = ws.cell(row=current_row, column=1, value=desc)
				c_desc.font = bold_font if (is_subtotal or is_grand_total) else regular_font
				c_desc.border = thin_border
				if is_grand_total:
					c_desc.fill = total_fill
				elif is_subtotal:
					c_desc.fill = sub_header_fill

				c_val = ws.cell(row=current_row, column=2, value=float(val or 0.0))
				c_val.font = bold_font if (is_subtotal or is_grand_total) else regular_font
				c_val.number_format = "#,##0.00"
				c_val.border = thin_border
				c_val.alignment = Alignment(horizontal="right")
				if is_grand_total:
					c_val.fill = total_fill
				elif is_subtotal:
					c_val.fill = sub_header_fill

				current_row += 1

			current_row += 2

			# 2. Payments Received Schedule
			ws.cell(row=current_row, column=1, value="2. PAYMENTS RECEIVED").font = section_font
			current_row += 1

			sec2_headers = ["Date", "Payment Reference", "Description / Allocation", "Amount Received"]
			for col_idx, h in enumerate(sec2_headers, start=1):
				c = ws.cell(row=current_row, column=col_idx, value=h)
				c.font = header_font
				c.fill = header_fill
				c.border = thin_border
				c.alignment = Alignment(horizontal="center" if col_idx <= 2 else "right" if col_idx == 4 else "left")
			current_row += 1

			payments = rec_data.get("payments", [])
			if payments:
				for p in payments:
					p_date = str(p.get("posting_date") or "")
					p_ref = str(p.get("voucher_no") or "")
					p_rem = str(p.get("remarks") or "Advance / project payment")
					p_amt = float(p.get("amount") or 0.0)

					ws.cell(row=current_row, column=1, value=p_date).border = thin_border
					ws.cell(row=current_row, column=2, value=p_ref).border = thin_border
					ws.cell(row=current_row, column=3, value=p_rem).border = thin_border
					c_amt = ws.cell(row=current_row, column=4, value=p_amt)
					c_amt.border = thin_border
					c_amt.number_format = "#,##0.00"
					c_amt.alignment = Alignment(horizontal="right")
					current_row += 1

				# Total Payments
				ws.cell(row=current_row, column=1, value="Total Payments Received").font = bold_font
				ws.cell(row=current_row, column=1).fill = sub_header_fill
				ws.cell(row=current_row, column=1).border = thin_border
				ws.cell(row=current_row, column=2, value="").fill = sub_header_fill
				ws.cell(row=current_row, column=2).border = thin_border
				ws.cell(row=current_row, column=3, value="").fill = sub_header_fill
				ws.cell(row=current_row, column=3).border = thin_border

				c_tot = ws.cell(row=current_row, column=4, value=float(paid_amt or 0.0))
				c_tot.font = bold_font
				c_tot.fill = sub_header_fill
				c_tot.border = thin_border
				c_tot.number_format = "#,##0.00"
				c_tot.alignment = Alignment(horizontal="right")
				current_row += 1
			else:
				ws.cell(row=current_row, column=1, value="No direct cash payments recorded for this project.").font = regular_font
				current_row += 1

			current_row += 2

			# 3. Credit Note Reallocation Schedule
			ws.cell(row=current_row, column=1, value="3. CREDIT NOTE REALLOCATION SCHEDULE").font = section_font
			current_row += 1

			sec3_headers = ["Credit Note Ref", "Return Against", "Reason / Remarks", "Target Invoice Settled", "Settlement Link / JV", "Credit Amount"]
			for col_idx, h in enumerate(sec3_headers, start=1):
				c = ws.cell(row=current_row, column=col_idx, value=h)
				c.font = header_font
				c.fill = header_fill
				c.border = thin_border
				c.alignment = Alignment(horizontal="center" if col_idx <= 2 else "right" if col_idx == 6 else "left")
			current_row += 1

			cns = rec_data.get("credit_notes", [])
			if cns:
				for cn in cns:
					cn_ref = str(cn.get("voucher_no") or "")
					ret_ag = str(cn.get("return_against") or "")
					cn_rem = str(cn.get("remarks") or "")
					allocs = cn.get("allocations", [])
					if allocs:
						for a in allocs:
							ws.cell(row=current_row, column=1, value=cn_ref).border = thin_border
							ws.cell(row=current_row, column=2, value=ret_ag).border = thin_border
							ws.cell(row=current_row, column=3, value=cn_rem).border = thin_border
							ws.cell(row=current_row, column=4, value=str(a.get("target_invoice") or "")).border = thin_border
							ws.cell(row=current_row, column=5, value=str(a.get("jv") or "-")).border = thin_border
							c_amt = ws.cell(row=current_row, column=6, value=float(a.get("amount") or 0.0))
							c_amt.border = thin_border
							c_amt.number_format = "#,##0.00"
							c_amt.alignment = Alignment(horizontal="right")
							current_row += 1
					else:
						ws.cell(row=current_row, column=1, value=cn_ref).border = thin_border
						ws.cell(row=current_row, column=2, value=ret_ag).border = thin_border
						ws.cell(row=current_row, column=3, value=cn_rem).border = thin_border
						ws.cell(row=current_row, column=4, value="Unallocated Open Credit").border = thin_border
						ws.cell(row=current_row, column=5, value="-").border = thin_border
						c_amt = ws.cell(row=current_row, column=6, value=float(cn.get("amount") or 0.0))
						c_amt.border = thin_border
						c_amt.number_format = "#,##0.00"
						c_amt.alignment = Alignment(horizontal="right")
						current_row += 1

				# Total CNs
				ws.cell(row=current_row, column=1, value="Total Credit Notes Issued").font = bold_font
				ws.cell(row=current_row, column=1).fill = sub_header_fill
				ws.cell(row=current_row, column=1).border = thin_border
				for ci in range(2, 6):
					ws.cell(row=current_row, column=ci, value="").fill = sub_header_fill
					ws.cell(row=current_row, column=ci).border = thin_border

				c_tot = ws.cell(row=current_row, column=6, value=float(cn_amt or 0.0))
				c_tot.font = bold_font
				c_tot.fill = sub_header_fill
				c_tot.border = thin_border
				c_tot.number_format = "#,##0.00"
				c_tot.alignment = Alignment(horizontal="right")
				current_row += 1
			else:
				ws.cell(row=current_row, column=1, value="No credit notes or return invoices recorded for this project.").font = regular_font
				current_row += 1

			current_row += 2

			# 4. Current Outstanding Invoices Schedule
			ws.cell(row=current_row, column=1, value="4. CURRENT OUTSTANDING INVOICES").font = section_font
			current_row += 1

			sec4_headers = ["Posting Date", "Due Date", "Invoice Reference", "Gross Invoiced", "Paid / Credited", "Balance Due"]
			for col_idx, h in enumerate(sec4_headers, start=1):
				c = ws.cell(row=current_row, column=col_idx, value=h)
				c.font = header_font
				c.fill = header_fill
				c.border = thin_border
				c.alignment = Alignment(horizontal="center" if col_idx <= 2 else "right" if col_idx >= 4 else "left")
			current_row += 1

			open_invs = rec_data.get("outstanding_invoices", [])
			if open_invs:
				for oi in open_invs:
					oi_date = str(oi.get("posting_date") or "")
					oi_due = str(oi.get("due_date") or "")
					oi_ref = str(oi.get("voucher_no") or "")
					oi_inv = float(oi.get("invoiced") or 0.0)
					oi_settled = float(oi.get("paid") or (oi_inv - float(oi.get("outstanding") or 0.0)))
					oi_out = float(oi.get("outstanding") or 0.0)

					ws.cell(row=current_row, column=1, value=oi_date).border = thin_border
					ws.cell(row=current_row, column=2, value=oi_due).border = thin_border
					ws.cell(row=current_row, column=3, value=oi_ref).border = thin_border

					c4 = ws.cell(row=current_row, column=4, value=oi_inv)
					c4.border = thin_border
					c4.number_format = "#,##0.00"
					c4.alignment = Alignment(horizontal="right")

					c5 = ws.cell(row=current_row, column=5, value=oi_settled)
					c5.border = thin_border
					c5.number_format = "#,##0.00"
					c5.alignment = Alignment(horizontal="right")

					c6 = ws.cell(row=current_row, column=6, value=oi_out)
					c6.font = bold_font
					c6.border = thin_border
					c6.number_format = "#,##0.00"
					c6.alignment = Alignment(horizontal="right")
					current_row += 1

				# Total Outstanding
				ws.cell(row=current_row, column=1, value="TOTAL REMAINING BALANCE DUE").font = bold_font
				ws.cell(row=current_row, column=1).fill = total_fill
				ws.cell(row=current_row, column=1).border = thin_border
				for ci in range(2, 6):
					ws.cell(row=current_row, column=ci, value="").fill = total_fill
					ws.cell(row=current_row, column=ci).border = thin_border

				c_tot = ws.cell(row=current_row, column=6, value=float(out_amt or 0.0))
				c_tot.font = bold_font
				c_tot.fill = total_fill
				c_tot.border = thin_border
				c_tot.number_format = "#,##0.00"
				c_tot.alignment = Alignment(horizontal="right")
				current_row += 1
			else:
				ws.cell(row=current_row, column=1, value="All invoices for this project are fully settled. Net balance due is 0.00.").font = regular_font
				current_row += 1

		else:
			# Executive Summary Cards for GL and standard AR
			if summary:
				ws.cell(row=current_row, column=1, value="EXECUTIVE SUMMARY").font = section_font
				current_row += 1

				if doc.report == "General Ledger":
					kpis = [
						("Opening Balance", summary.get("opening_balance", 0)),
						("Total Invoiced (Debits)", summary.get("total_invoiced", 0)),
						("Credit Notes & Returns", summary.get("total_credit_notes", 0)),
						("Payments / Receipts", summary.get("total_payments", 0)),
						("Closing Balance Due", summary.get("closing_balance", 0)),
					]
				else:
					gross_inv = summary.get("gross_invoiced") or summary.get("total_invoiced", 0)
					cn_amt = summary.get("total_credit_notes", 0)
					net_inv = summary.get("net_invoiced") or (gross_inv - cn_amt)
					kpis = [
						("Total Invoiced (Gross)", gross_inv),
						("Credit Notes & Returns", -cn_amt if cn_amt else 0),
						("Net Billed Amount", net_inv),
						("Total Paid Amount", summary.get("total_paid", 0)),
						("Net Outstanding Due", summary.get("total_outstanding", 0)),
					]

				for col_idx, (label, val) in enumerate(kpis, start=1):
					lbl_cell = ws.cell(row=current_row, column=col_idx, value=label)
					lbl_cell.font = bold_font
					lbl_cell.fill = sub_header_fill
					lbl_cell.border = thin_border
					lbl_cell.alignment = Alignment(horizontal="center")

					val_cell = ws.cell(row=current_row + 1, column=col_idx, value=float(val or 0.0))
					val_cell.font = bold_font
					val_cell.number_format = "#,##0.00"
					val_cell.fill = card_fill
					val_cell.border = thin_border
					val_cell.alignment = Alignment(horizontal="right")

				current_row += 3

			# Table Columns
			if doc.report == "General Ledger":
				headers = [
					"Date",
					"Type",
					"Reference",
					"Particulars / Allocation Details",
					"Project",
					"Debit",
					"Credit",
					"Balance",
				]
				for col_idx, header in enumerate(headers, start=1):
					cell = ws.cell(row=current_row, column=col_idx, value=header)
					cell.font = header_font
					cell.fill = header_fill
					cell.alignment = Alignment(horizontal="center" if col_idx <= 5 else "right")
					cell.border = thin_border
			else:
				headers = [
					"Date",
					"Due Date",
					"Reference",
					"Project & Details",
					"Invoiced",
					"Paid",
					"Credit Note",
					"Outstanding",
				]
				for col_idx, header in enumerate(headers, start=1):
					cell = ws.cell(row=current_row, column=col_idx, value=header)
					cell.font = header_font
					cell.fill = header_fill
					cell.alignment = Alignment(horizontal="center" if col_idx <= 4 else "right")
					cell.border = thin_border

			current_row += 1

			# Table Data
			if doc.report == "General Ledger":
				for row in res:
					r_date = row.get("posting_date")
					r_vtype = row.get("voucher_type") or ""
					r_vno = row.get("voucher_no") or ""
					r_account = row.get("account") or ""
					r_project = row.get("project") or ""
					r_remarks = row.get("remarks") or ""
					r_debit = row.get("debit")
					r_credit = row.get("credit")
					r_balance = row.get("balance")

					details = []
					if row.get("is_return") and row.get("return_against"):
						details.append(f"[CREDIT NOTE - Return against {row.get('return_against')}]")
					if row.get("allocations"):
						alloc_str = ", ".join(
							[
								f"{a.reference_name} ({frappe.format_value(a.allocated_amount, 'Currency')})"
								for a in row.get("allocations")
							]
						)
						details.append(f"Allocated to: {alloc_str}")
					if doc.show_remarks and r_remarks:
						rem = str(r_remarks).strip()
						if rem and rem not in ["No Remarks", "No Remarks.", "None", r_vno]:
							details.append(rem)
					if not r_date and r_account:
						details.append(r_account)

					detail_str = " | ".join(details) if details else ""
					is_total_row = not r_date and (
						"Total" in str(r_account)
						or "Closing" in str(r_account)
						or "Opening" in str(r_account)
					)

					row_data = [
						str(r_date) if r_date else "",
						r_vtype,
						r_vno,
						detail_str or r_account,
						r_project,
						float(r_debit) if r_debit else 0.0,
						float(r_credit) if r_credit else 0.0,
						float(r_balance) if r_balance is not None else 0.0,
					]

					for col_idx, val in enumerate(row_data, start=1):
						c = ws.cell(row=current_row, column=col_idx, value=val)
						c.font = bold_font if is_total_row else regular_font
						c.border = thin_border
						if is_total_row:
							c.fill = total_fill
						if col_idx >= 6:
							c.number_format = "#,##0.00"
							c.alignment = Alignment(horizontal="right")

					current_row += 1
			else:
				for row in res:
					r_date = row.get("posting_date")
					r_due = row.get("due_date")
					r_vno = row.get("voucher_no") or ""
					if row.get("customer_ref"):
						r_vno = f"{r_vno} (Ref: {row.get('customer_ref')})"

					r_project = row.get("project") or ""
					details = []
					if row.get("is_payment_entry") or row.get("status") == "Unallocated Payment":
						if row.get("mode_of_payment"):
							details.append(f"Mode: {row.get('mode_of_payment')}")
						if row.get("pe_settled_invoices"):
							tot_applied = sum(si['amount'] for si in row.get("pe_settled_invoices"))
							inv_list = ", ".join(si['invoice'] for si in row.get("pe_settled_invoices"))
							details.append(f"Applied: {tot_applied:,.2f} across {len(row.get('pe_settled_invoices'))} invoice(s) ({inv_list})")
							details.append("Remaining unallocated credit balance available on account")
						else:
							details.append("Advance payment credit available for allocation")
					else:
						# Invoice settlement details matching PDF
						if row.get("allocations"):
							for alloc in row.get("allocations"):
								amt_fmt = f"{alloc.get('amount', 0):,.2f}"
								ref_info = ""
								if alloc.get("reference_no"):
									ref_info = f" (Ref: {alloc.get('reference_no')}"
									if alloc.get("mode_of_payment"):
										ref_info += f" • {alloc.get('mode_of_payment')}"
									ref_info += ")"
								elif alloc.get("mode_of_payment"):
									ref_info = f" ({alloc.get('mode_of_payment')})"

								if alloc.get("is_credit_note"):
									cn_ref = alloc.get("credit_note_ref") or alloc.get("voucher_no")
									details.append(f"Credit Note Applied: {cn_ref} — {amt_fmt}")
								elif alloc.get("is_advance"):
									details.append(f"Settled via Advance: {alloc.get('voucher_no')}{ref_info} — {amt_fmt}")
								else:
									details.append(f"Settled via Payment: {alloc.get('voucher_no')}{ref_info} — {amt_fmt}")

						if row.get("is_return") and row.get("return_against"):
							details.append(f"[Return against {row.get('return_against')}]")
						if doc.show_remarks:
							rem = str(row.get("invoice_remarks") or row.get("remarks") or "").strip()
							if rem and rem not in ["No Remarks", "No Remarks.", "None", row.get("voucher_no")]:
								details.append(rem)

					# Prepend project if present so it mirrors "Project & Details"
					if r_project:
						details.insert(0, f"Project: {r_project}")

					detail_str = " | ".join(details) if details else ""
					r_invoiced = float(row.get("invoiced") or row.get("invoiced_amount") or 0.0)
					r_paid = float(row.get("paid") or row.get("paid_amount") or 0.0)
					r_credit = float(row.get("credit_note") or row.get("credit_note_amount") or 0.0)
					r_outstanding = float(row.get("outstanding") or row.get("outstanding_amount") or 0.0)

					row_data = [
						str(r_date) if r_date else "",
						str(r_due) if r_due else "",
						r_vno,
						detail_str,
						r_invoiced,
						r_paid,
						r_credit,
						r_outstanding,
					]

					for col_idx, val in enumerate(row_data, start=1):
						c = ws.cell(row=current_row, column=col_idx, value=val)
						c.font = regular_font
						c.border = thin_border
						if 5 <= col_idx <= 8:
							c.number_format = "#,##0.00"
							c.alignment = Alignment(horizontal="right")
						elif col_idx <= 2:
							c.alignment = Alignment(horizontal="center")

					current_row += 1

		# Ageing Section
		if ageing and ageing[0]:
			current_row += 2
			ws.cell(
				row=current_row,
				column=1,
				value=f"Ageing Summary (Based on {ageing[0].get('ageing_based_on', 'Due Date')})",
			).font = section_font
			current_row += 1
			ageing_headers = [
				"0 - 30 Days",
				"30 - 60 Days",
				"60 - 90 Days",
				"90 - 120 Days",
				"Above 120 Days",
				"Total Outstanding",
			]
			ageing_vals = [
				ageing[0].get("range1", 0),
				ageing[0].get("range2", 0),
				ageing[0].get("range3", 0),
				ageing[0].get("range4", 0),
				ageing[0].get("range5", 0),
				ageing[0].get("total_outstanding")
				or sum([ageing[0].get(f"range{i}", 0) for i in range(1, 6)]),
			]
			for col_idx, h in enumerate(ageing_headers, start=1):
				c = ws.cell(row=current_row, column=col_idx, value=h)
				c.font = bold_font
				c.fill = sub_header_fill
				c.border = thin_border
				c.alignment = Alignment(horizontal="center")

			current_row += 1
			for col_idx, v in enumerate(ageing_vals, start=1):
				c = ws.cell(row=current_row, column=col_idx, value=float(v or 0))
				c.font = bold_font
				c.number_format = "#,##0.00"
				c.alignment = Alignment(horizontal="right")
				c.border = thin_border

		# Auto-adjust column widths
		for col_cells in ws.columns:
			max_len = max(len(str(cell.value or "")) for cell in col_cells)
			col_letter = get_column_letter(col_cells[0].column)
			ws.column_dimensions[col_letter].width = max(min(max_len + 3, 55), 12)

	if sheets_created == 0:
		return None

	output = io.BytesIO()
	wb.save(output)
	output.seek(0)
	return output


@frappe.whitelist()
def send_emails(document_name, from_scheduler=False, posting_date=None):
	doc = frappe.get_doc("Process Statement Of Accounts", document_name)
	report = get_report_pdf(doc, consolidated=False)

	if report:
		for customer, report_pdf in report.items():
			context = get_context(customer, doc)
			filename = frappe.render_template(doc.pdf_name, context)
			attachments = [{"fname": filename + ".pdf", "fcontent": report_pdf}]

			recipients, cc = get_recipients_and_cc(customer, doc)
			if not recipients:
				continue

			subject = frappe.render_template(doc.subject, context)
			message = frappe.render_template(doc.body, context)

			if doc.sender:
				sender_email = frappe.db.get_value("Email Account", doc.sender, "email_id")
			else:
				sender_email = frappe.session.user

			frappe.enqueue(
				queue="short",
				method=frappe.sendmail,
				recipients=recipients,
				sender=sender_email,
				cc=cc,
				subject=subject,
				message=message,
				now=True,
				reference_doctype="Process Statement Of Accounts",
				reference_name=document_name,
				attachments=attachments,
			)

		if doc.enable_auto_email and from_scheduler:
			new_to_date = getdate(posting_date or today())
			if doc.frequency == "Weekly":
				new_to_date = add_days(new_to_date, 7)
			else:
				new_to_date = add_months(new_to_date, 1 if doc.frequency == "Monthly" else 3)
			new_from_date = add_months(new_to_date, -1 * doc.filter_duration)
			doc.add_comment(
				"Comment", "Emails sent on: " + frappe.utils.format_datetime(frappe.utils.now())
			)
			if doc.report == "General Ledger":
				doc.db_set("to_date", new_to_date, commit=True)
				doc.db_set("from_date", new_from_date, commit=True)
			else:
				doc.db_set("posting_date", new_to_date, commit=True)
		return True
	else:
		return False


@frappe.whitelist()
def send_auto_email():
	selected = frappe.get_list(
		"Process Statement Of Accounts",
		filters={"enable_auto_email": 1},
		or_filters={"to_date": today(), "posting_date": today()},
	)
	for entry in selected:
		send_emails(entry.name, from_scheduler=True)
	return True
